"""PowerPoint (PPTX) adapter preserving slide hierarchy, tables, and speaker notes."""

from pathlib import Path
from typing import List, Any

from .base import BaseAdapter, AdapterPage, AdapterResult
from ..config import IngestionConfig
from ..schemas import WarningInfo


def _extract_shape_text(shape: Any) -> List[str]:
    """Recursively extract text fragments from shapes, grouped shapes, and tables."""
    texts: List[str] = []

    # 1. Grouped shapes recursion
    if hasattr(shape, "shapes"):
        for sub_shape in shape.shapes:
            texts.extend(_extract_shape_text(sub_shape))
        return texts

    # 2. Table extraction
    if shape.has_table:
        table_rows = []
        for row in shape.table.rows:
            row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            deduped = []
            for cell in row_cells:
                if not deduped or cell != deduped[-1]:
                    deduped.append(cell)
            table_rows.append(" | ".join(deduped))
        if table_rows:
            texts.append("\n".join(table_rows))
        return texts

    # 3. Standard Text Frame
    if shape.has_text_frame:
        tf_text = shape.text_frame.text.strip()
        if tf_text:
            texts.append(tf_text)

    return texts


class PPTXAdapter(BaseAdapter):
    """Layout-aware PowerPoint extractor preserving slide sequence, titles, and notes."""

    def extract(self, path: Path, config: IngestionConfig) -> AdapterResult:
        from pptx import Presentation

        prs = Presentation(str(path))
        num_slides = len(prs.slides)
        warnings: List[WarningInfo] = []

        if num_slides > config.max_pages:
            return AdapterResult(
                pages=[],
                warnings=[
                    WarningInfo(
                        code="ERR_PAGE_LIMIT_EXCEEDED",
                        message=f"Presentation slide count ({num_slides}) exceeds maximum allowed limit ({config.max_pages})",
                    )
                ],
                source_type="pptx",
            )

        pages: List[AdapterPage] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_elements: List[str] = []

            # 1. Title first
            title_text = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text_frame.text.strip()
                if title_text:
                    slide_elements.append(title_text)

            # 2. Body shapes sorted by spatial position (top-to-bottom, left-to-right)
            other_shapes = [s for s in slide.shapes if s != slide.shapes.title]
            other_shapes.sort(key=lambda s: (getattr(s, "top", 0), getattr(s, "left", 0)))

            for shape in other_shapes:
                shape_texts = _extract_shape_text(shape)
                for t in shape_texts:
                    if t and t != title_text:
                        slide_elements.append(t)

            # 3. Speaker notes labeled distinctly
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_elements.append(f"--- Speaker Notes ---\n{notes}")

            slide_full_text = "\n\n".join(slide_elements).strip()

            pages.append(
                AdapterPage(
                    page_num=slide_num,
                    label="slide",
                    text=slide_full_text,
                    method="native",
                    ocr_confidence=None,
                    warnings=[],
                )
            )

        return AdapterResult(pages=pages, warnings=warnings, source_type="pptx")
