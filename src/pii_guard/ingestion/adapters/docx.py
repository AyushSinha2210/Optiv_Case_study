"""Structure-preserving Word document (DOCX) adapter interleaving paragraphs and tables."""

from pathlib import Path
from typing import List

from .base import BaseAdapter, AdapterPage, AdapterResult
from ..config import IngestionConfig
from ..schemas import WarningInfo


class DOCXAdapter(BaseAdapter):
    """Layout-aware DOCX extractor preserving true visual document flow."""

    def extract(self, path: Path, config: IngestionConfig) -> AdapterResult:
        import docx
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        doc = docx.Document(str(path))
        warnings: List[WarningInfo] = []
        body_elements: List[str] = []

        # 1. Header extraction
        header_texts: List[str] = []
        for section in doc.sections:
            if section.header and section.header.paragraphs:
                for p in section.header.paragraphs:
                    t = p.text.strip()
                    if t and t not in header_texts:
                        header_texts.append(t)

        if header_texts:
            body_elements.append("--- Header ---\n" + "\n".join(header_texts))

        # 2. Iterate document body in true sequential XML child order
        for child in doc._element.body:
            tag = child.tag
            if tag.endswith("p"):
                para = Paragraph(child, doc)
                text = para.text.strip()
                if text:
                    body_elements.append(text)
            elif tag.endswith("tbl"):
                tbl = Table(child, doc)
                table_rows: List[str] = []
                for row in tbl.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    # Deduplicate merged cells that share identical text consecutively
                    deduped_cells = []
                    for c in cells:
                        if not deduped_cells or c != deduped_cells[-1]:
                            deduped_cells.append(c)
                    table_rows.append(" | ".join(deduped_cells))
                if table_rows:
                    body_elements.append("\n".join(table_rows))

        # 3. Footer extraction
        footer_texts: List[str] = []
        for section in doc.sections:
            if section.footer and section.footer.paragraphs:
                for p in section.footer.paragraphs:
                    t = p.text.strip()
                    if t and t not in footer_texts:
                        footer_texts.append(t)

        if footer_texts:
            body_elements.append("--- Footer ---\n" + "\n".join(footer_texts))

        full_text = "\n\n".join(body_elements).strip()

        pages = [
            AdapterPage(
                page_num=1,
                label="page",
                text=full_text,
                method="native",
                ocr_confidence=None,
                warnings=[],
            )
        ]

        return AdapterResult(pages=pages, warnings=warnings, source_type="docx")
